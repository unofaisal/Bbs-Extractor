# Copyright (c) 2026, one and contributors
# For license information, please see license.txt


import os
import json
from io import BytesIO

 
import frappe
from frappe import _
from frappe.model.document import Document

from agent_builder.native_api.agent.runner import SessionProvenance, run_headless_agent



class BBSSchedule(Document):
    pass


# ---------------------------------------------------------------------------
# Excel Export
# ---------------------------------------------------------------------------

def get_export_columns(standard: str) -> list:
    """Return column layout based on the BBS standard."""
    cols = [
        ("member_group", "Member Group", 18),
        ("member", "Member", 18),
        ("bar_mark", "Bar Mark", 10),
        ("bar_size", "Bar Size", 9),
        ("no_of_members", "No. of Mbrs", 9),
        ("no_of_bars_each", "Bars Each", 9),
        ("total_no", "Total No.", 9),
    ]
    
    if standard == "BS 8666 (Shape Code)":
        cols.extend([
            ("length_mm", "Length (mm)", 12),
            ("shape_code", "Shape Code", 10),
            ("dim_a_mm", "A (mm)", 8),
            ("dim_b_mm", "B (mm)", 8),
            ("dim_c_mm", "C (mm)", 8),
            ("dim_d_mm", "D (mm)", 8),
            ("dim_e_r_mm", "E/R (mm)", 8),
        ])
    elif standard == "BS 8666 (Spacing)":
        cols.extend([
            ("bar_spacing_mm", "Spacing (mm)", 12),
            ("length_mm", "Length (mm)", 12),
        ])
    elif standard == "BS 4466":
        cols.extend([
            ("length_m", "Length (m)", 12),
        ])
    else:
        # Fallback: show everything if standard is missing
        cols.extend([
            ("length_mm", "Length (mm)", 12),
            ("length_m", "Length (m)", 12),
            ("bar_spacing_mm", "Spacing (mm)", 12),
            ("shape_code", "Shape Code", 10),
            ("dim_a_mm", "A (mm)", 8),
            ("dim_b_mm", "B (mm)", 8),
            ("dim_c_mm", "C (mm)", 8),
            ("dim_d_mm", "D (mm)", 8),
            ("dim_e_r_mm", "E/R (mm)", 8),
        ])

    cols.extend([
        ("weight_kg", "Weight (kg)", 12),
        ("review", "Review / Notes", 25),
        ("mapped_product_code", "Mapped Code", 18),
        ("mapped_product_name", "Mapped Product", 20),
    ])
    
    return cols


def group_items_hierarchy(items: list) -> list:
    """Group items into a 2-level hierarchy: [ {group, members: [ {member, items: []} ] } ]"""
    hierarchy = []
    for item in items:
        grp = item.member_group or ""
        mem = item.member or ""
        
        if not hierarchy or hierarchy[-1]["group"] != grp:
            hierarchy.append({"group": grp, "members": []})
            
        current_group = hierarchy[-1]
        if not current_group["members"] or current_group["members"][-1]["member"] != mem:
            current_group["members"].append({"member": mem, "items": []})
            
        current_group["members"][-1]["items"].append(item)
        
    return hierarchy


@frappe.whitelist()
def get_excel(name: str):
    """Stream a clean, formatted .xlsx for a single BBS Schedule."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from io import BytesIO # Ensure this is imported at the top of your file

    doc = frappe.get_doc("BBS Schedule", name)
    columns = get_export_columns(doc.standard)
    total_cols = len(columns)

    wb = Workbook()
    ws = wb.active
    ws.title = "BBS Schedule"

    # --- Styles ---
    title_font = Font(bold=True, size=14)
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9D9D9") # Light grey for headers
    group_fill = PatternFill("solid", fgColor="F2F2F2")  # Very light grey for groups
    member_fill = PatternFill("solid", fgColor="FFFFFF") # White for members
    
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # --- Header Block ---
    row = 1
    ws.cell(row=row, column=1, value=doc.project_name or doc.schedule_name).font = title_font
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
    row += 1

    meta = [
        ("Schedule Ref.", doc.schedule_ref),
        ("Standard", doc.standard),
        ("Status", doc.status),
        ("Total Weight (kg)", doc.total_weight_kg),
    ]
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1 # Spacer
    table_start_row = row

    # --- Table Header ---
    for i, (_, label, _) in enumerate(columns):
        cell = ws.cell(row=row, column=i + 1, value=label)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[row].height = 30
    row += 1

    # --- Data Rows (Hierarchical Grouping) ---
    hierarchy = group_items_hierarchy(doc.items)
    text_fields = {"member_group", "member", "bar_mark", "bar_size", "shape_code", "mapped_product_code", "mapped_product_name", "review"}

    for group_data in hierarchy:
        group_start_row = row
        
        for member_data in group_data["members"]:
            member_start_row = row
            
            for item in member_data["items"]:
                for i, (fieldname, _, _) in enumerate(columns):
                    # Leave col 1 and 2 blank for merged cells, handle below
                    if fieldname in ("member_group", "member"):
                        continue
                        
                    val = item.get(fieldname)
                    cell = ws.cell(row=row, column=i + 1, value=val)
                    cell.border = border
                    cell.alignment = left_align if fieldname in text_fields else center
                row += 1
                
            member_end_row = row - 1
            
            # Merge Member cell (Column 2)
            member_cell = ws.cell(row=member_start_row, column=2, value=member_data["member"])
            member_cell.font = bold
            member_cell.fill = member_fill
            member_cell.alignment = center
            member_cell.border = border
            
            if member_end_row > member_start_row:
                ws.merge_cells(start_row=member_start_row, start_column=2, end_row=member_end_row, end_column=2)
                # Re-apply border to all cells in the merged range
                for r in range(member_start_row, member_end_row + 1):
                    ws.cell(row=r, column=2).border = border
                    ws.cell(row=r, column=2).fill = member_fill

        group_end_row = row - 1
        
        # Merge Member Group cell (Column 1)
        if group_data["group"]:
            group_cell = ws.cell(row=group_start_row, column=1, value=group_data["group"])
            group_cell.font = bold
            group_cell.fill = group_fill
            group_cell.alignment = center
            group_cell.border = border
            
            if group_end_row > group_start_row:
                ws.merge_cells(start_row=group_start_row, start_column=1, end_row=group_end_row, end_column=1)
                # Re-apply border to all cells in the merged range
                for r in range(group_start_row, group_end_row + 1):
                    ws.cell(row=r, column=1).border = border
                    ws.cell(row=r, column=1).fill = group_fill

    # --- Column Widths ---
    for i, (_, _, width) in enumerate(columns):
        ws.column_dimensions[get_column_letter(i + 1)].width = width

    # Freeze panes below header and to the right of the Member columns
    ws.freeze_panes = ws.cell(row=table_start_row + 1, column=3)

    # --- Output ---
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    frappe.response["filename"] = f"{doc.schedule_ref or doc.schedule_name}.xlsx"
    frappe.response["filecontent"] = buf.getvalue()
    
    # Changed from "binary" to "download" 
    # This tells Frappe to send Content-Disposition: attachment
    frappe.response["type"] = "download"
 
 
def _resolve_customer_email(customer: str | None) -> str:
    """Primary email off the Customer's linked Contact, if one exists.
    Returns "" (not None) so the frontend can drop it straight into the
    recipients field either way — blank means "let the user fill it in"."""
    if not customer:
        return ""
    try:
        contact_name = frappe.db.get_value(
            "Dynamic Link",
            {"link_doctype": "Customer", "link_name": customer, "parenttype": "Contact"},
            "parent",
        )
        if not contact_name:
            return ""
        email = frappe.db.get_value("Contact", contact_name, "email_id")
        return email or ""
    except Exception:
        frappe.log_error("bbs_extractor: customer email lookup failed", frappe.get_traceback())
        return ""
 
 
def _build_doc_context(doc, review_items) -> str:
    lines = [
        f"Schedule: {doc.schedule_name} (Ref: {doc.schedule_ref or '-'})",
        f"Project: {doc.project_name or '-'}",
        f"Customer: {doc.customer or '-'}",
        "",
        "Items flagged for review:",
    ]
    for item in review_items:
        lines.append(
            f"- Bar Mark {item.bar_mark or '-'} "
            f"(Member: {item.member or '-'}, Size: {item.bar_size or '-'}): {item.review}"
        )
    return "\n".join(lines)
 
 
@frappe.whitelist()
def get_customer_email(customer: str) -> str:
    """Thin whitelisted wrapper so the frontend can fetch the customer's
    email on its own — used by the plain 'Create Email' button, which
    doesn't run an agent at all."""
    return _resolve_customer_email(customer)
 
 
@frappe.whitelist()
def draft_clarification_email(name: str) -> dict:
    """AI-drafted clarification email for a BBS Schedule's flagged items,
    run through the standard headless agent pipeline (same as triggers)
    so it's stored as a normal Agent session.
 
    Returns {subject, message, recipients, session_id} ready to hand
    straight to frappe.views.CommunicationComposer on the frontend.
    """
    doc = frappe.get_doc("BBS Schedule", name)
    review_items = [i for i in doc.items if (i.review or "").strip()]
 
    if not review_items:
        frappe.throw(_("No items on this schedule have a Clarify/Review note."))
 
    doc_context = _build_doc_context(doc, review_items)
    recipient_email = _resolve_customer_email(doc.customer)
 
    provenance = SessionProvenance(
        trigger_type="Manual",
        trigger_source="BBS Schedule",
        trigger_ref=name,
    )
    agent = frappe.get_doc("Skill", "email_drafter_instructions").get("content")
 
    result = run_headless_agent(
        agent_name=None,
        input_message=doc_context,
        provenance=provenance,
        user=frappe.session.user,
        skill_injection=agent,
    )
 
    if result.ended_reason != "Completed":
        frappe.log_error(
            f"bbs_extractor: clarification email draft ended with {result.ended_reason}",
            "bbs_extractor email draft",
        )
        frappe.throw(_("Couldn't draft the email automatically. Please use 'Create Email' instead."))
 
    raw = (result.response or "").strip()
    if raw.startswith("```"):
        raw = raw.strip("`")
        if raw.lower().startswith("json"):
            raw = raw[4:].strip()
 
    try:
        # strict=False: models reliably put literal newlines inside the
        # body string (for paragraph breaks) instead of escaping them as
        # \n. That's invalid strict JSON but perfectly recoverable — no
        # ambiguity in intent — so tolerate it rather than rejecting a
        # good draft over one escaping slip.
        draft = json.loads(raw, strict=False)
    except Exception:
        frappe.log_error(
            f"bbs_extractor: clarification email draft unparseable — session {result.session_id}",
            frappe.get_traceback(),
        )
        frappe.throw(_("AI draft came back in an unexpected format. Please use 'Create Email' instead."))
 
    if not isinstance(draft, dict) or "subject" not in draft or "body" not in draft:
        frappe.throw(_("AI draft came back in an unexpected format. Please use 'Create Email' instead."))
 
    return {
        "subject": draft["subject"],
        "message": draft["body"],
        "recipients": recipient_email,
        "session_id": result.session_id,
    }